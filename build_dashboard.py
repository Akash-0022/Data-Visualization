"""
Builds the Excel dashboard from the cleaned dataset.

Run from inside the "Data Visualization" folder:
    python build_dashboard.py

Output: Exoplanet_Dashboard.xlsx  with 3 sheets
    Cleaned Data -> the cleaned dataset
    Analysis     -> the small tables used by the charts
    Dashboard    -> KPI cards + charts
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.line import LineProperties

CLEAN_FILE = "../Data Cleaning/cleaned_exoplanet_data.csv"
OUT_FILE = "Exoplanet_Dashboard.xlsx"

df = pd.read_csv(CLEAN_FILE)

# ---------- colours / styles used everywhere ----------
DARK = "1F3864"
BLUE = "2E75B6"
LIGHT = "D9E2F3"
GREY = "F2F2F2"

white_bold = Font(color="FFFFFF", bold=True, size=11)
title_font = Font(color="FFFFFF", bold=True, size=18)
thin = Side(style="thin", color="B4C6E7")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ======================================================================
# SHEET 1 : CLEANED DATA
# ======================================================================
ws1 = wb.active
ws1.title = "Cleaned Data"

for row in dataframe_to_rows(df, index=False, header=True):
    ws1.append(row)

for cell in ws1[1]:
    cell.fill = PatternFill("solid", fgColor=DARK)
    cell.font = white_bold
    cell.alignment = Alignment(horizontal="center", vertical="center")

ws1.freeze_panes = "A2"
ws1.auto_filter.ref = ws1.dimensions
ws1.row_dimensions[1].height = 22

widths = {"A": 26, "B": 14, "C": 10, "D": 15, "E": 15, "F": 11, "G": 20,
          "H": 20, "I": 10, "J": 17, "K": 11, "L": 15, "M": 20}
for col, w in widths.items():
    ws1.column_dimensions[col].width = w

for row in ws1.iter_rows(min_row=2, min_col=3, max_col=11):
    for cell in row:
        cell.number_format = "0.0000"

# ======================================================================
# SHEET 2 : ANALYSIS TABLES (these feed the charts)
# ======================================================================
ws2 = wb.create_sheet("Analysis")

def put_table(ws, top_row, col, heading, table):
    """writes a small table with a heading and returns the last row used"""
    c = ws.cell(row=top_row, column=col, value=heading)
    c.font = Font(bold=True, size=12, color=DARK)

    r = top_row + 1
    for j, name in enumerate(table.columns):
        h = ws.cell(row=r, column=col + j, value=name)
        h.fill = PatternFill("solid", fgColor=BLUE)
        h.font = white_bold
        h.alignment = Alignment(horizontal="center")
        h.border = box
    for _, values in table.iterrows():
        r += 1
        for j, v in enumerate(values):
            cell = ws.cell(row=r, column=col + j, value=v)
            cell.border = box
            if isinstance(v, float):
                cell.number_format = "0.0000"
    return r

# Table 1 - planets by status
status_tbl = (df["planet_status"].value_counts()
              .rename_axis("Planet Status").reset_index(name="Number of Planets"))
end1 = put_table(ws2, 1, 1, "Table 1 - Number of Planets by Planet Status", status_tbl)

# Table 2 - top 10 planets by mass
top10 = (df.nlargest(10, "mass_mj")[["name", "mass_mj"]]
         .rename(columns={"name": "Planet Name", "mass_mj": "Mass (MJup)"}))
end2 = put_table(ws2, end1 + 3, 1, "Table 2 - Top 10 Planets by Mass", top10)
# the chart labels copy the format of these cells, 4 decimals is too long
for r in range(end1 + 5, end2 + 1):
    ws2.cell(row=r, column=2).number_format = "0.000"

# Table 3 - planets by mass category
order = ["Small (< 0.1 MJ)", "Medium (0.1 - 1 MJ)", "Large (1 - 4 MJ)",
         "Very Large (> 4 MJ)", "Unknown"]
cat_tbl = (df["mass_category"].value_counts().reindex(order)
           .rename_axis("Mass Category").reset_index(name="Number of Planets"))
end3 = put_table(ws2, end2 + 3, 1, "Table 3 - Number of Planets by Mass Category", cat_tbl)

# Table 4 - summary statistics
mass = df["mass_mj"].dropna()
radius = df["radius"].dropna()
stat_tbl = pd.DataFrame({
    "Statistic": ["Count", "Mean", "Median", "Minimum", "Maximum",
                  "Standard Deviation", "Variance", "IQR"],
    "Mass (MJup)": [mass.count(), mass.mean(), mass.median(), mass.min(), mass.max(),
                    mass.std(), mass.var(), mass.quantile(.75) - mass.quantile(.25)],
    "Radius (RJup)": [radius.count(), radius.mean(), radius.median(), radius.min(),
                      radius.max(), radius.std(), radius.var(),
                      radius.quantile(.75) - radius.quantile(.25)],
})
put_table(ws2, end3 + 3, 1, "Table 4 - Summary Statistics", stat_tbl)

# Table 5 - mass vs radius (only rows where both are available) -> scatter plot
scatter_df = (df.dropna(subset=["mass_mj", "radius"])[["name", "mass_mj", "radius"]]
              .rename(columns={"name": "Planet Name", "mass_mj": "Mass (MJup)",
                               "radius": "Radius (RJup)"}))
scat_end = put_table(ws2, 1, 5, "Table 5 - Mass vs Radius (planets having both values)",
                     scatter_df)

for col, w in {"A": 24, "B": 20, "C": 16, "E": 24, "F": 14, "G": 14}.items():
    ws2.column_dimensions[col].width = w

# ======================================================================
# SHEET 3 : DASHBOARD
# ======================================================================
ws3 = wb.create_sheet("Dashboard")
ws3.sheet_view.showGridLines = False

ws3.column_dimensions["A"].width = 2
for i in range(2, 20):
    ws3.column_dimensions[get_column_letter(i)].width = 10.5

# ---------- title bar ----------
ws3.merge_cells("B2:O3")
t = ws3["B2"]
t.value = "EXOPLANET DATA ANALYSIS DASHBOARD"
t.font = title_font
t.fill = PatternFill("solid", fgColor=DARK)
t.alignment = Alignment(horizontal="center", vertical="center")
for row in ws3["B2:O3"]:
    for cell in row:
        cell.fill = PatternFill("solid", fgColor=DARK)
ws3.row_dimensions[2].height = 24
ws3.row_dimensions[3].height = 20

ws3.merge_cells("B4:O4")
s = ws3["B4"]
s.value = ("Source: Exoplanet catalogue (5,986 confirmed planets)   |   "
           "Mass in Jupiter masses (MJup), Radius in Jupiter radii (RJup)")
s.font = Font(italic=True, size=9, color="595959")
s.alignment = Alignment(horizontal="center")

# ---------- KPI cards ----------
# the values are written as Excel formulas so that the numbers are really
# calculated inside Excel and not just typed in
n = len(df) + 1              # last data row on the "Cleaned Data" sheet
DATA = "'Cleaned Data'"
mass_range = f"{DATA}!K2:K{n}"
rad_range = f"{DATA}!I2:I{n}"
name_range = f"{DATA}!A2:A{n}"

cards = [
    ("TOTAL PLANETS", f"=COUNTA({DATA}!A2:A{n})", "0"),
    ("HIGHEST MASS PLANET", f"=INDEX({name_range},MATCH(MAX({mass_range}),{mass_range},0))", "@"),
    ("LOWEST MASS PLANET", f"=INDEX({name_range},MATCH(MIN({mass_range}),{mass_range},0))", "@"),
    ("HIGHEST MASS (MJup)", f"=MAX({mass_range})", "0.000"),
    ("AVERAGE MASS (MJup)", f"=AVERAGE({mass_range})", "0.000"),
    ("AVERAGE RADIUS (RJup)", f"=AVERAGE({rad_range})", "0.000"),
]

positions = [(6, 2), (6, 7), (6, 12), (10, 2), (10, 7), (10, 12)]   # (row, col)

for (label, formula, fmt), (r, c) in zip(cards, positions):
    ws3.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 3)
    head = ws3.cell(row=r, column=c, value=label)
    head.font = Font(bold=True, size=9, color="FFFFFF")
    head.alignment = Alignment(horizontal="center", vertical="center")

    ws3.merge_cells(start_row=r + 1, start_column=c, end_row=r + 2, end_column=c + 3)
    val = ws3.cell(row=r + 1, column=c, value=formula)
    val.font = Font(bold=True, size=15, color=DARK)
    val.alignment = Alignment(horizontal="center", vertical="center")
    val.number_format = fmt

    for cc in range(c, c + 4):
        ws3.cell(row=r, column=cc).fill = PatternFill("solid", fgColor=BLUE)
        for rr in (r + 1, r + 2):
            cell = ws3.cell(row=rr, column=cc)
            cell.fill = PatternFill("solid", fgColor=GREY)
            cell.border = box

    ws3.row_dimensions[r].height = 16
    ws3.row_dimensions[r + 1].height = 18
    ws3.row_dimensions[r + 2].height = 12

def value_labels(chart, fmt=None):
    """show only the number on the bars, not the series / category name"""
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.dataLabels.showSerName = False
    chart.dataLabels.showCatName = False
    chart.dataLabels.showLegendKey = False
    chart.dataLabels.showPercent = False
    chart.dataLabels.showBubbleSize = False
    if fmt:
        chart.dataLabels.numFmt = fmt


def fix_chart(chart, colour=BLUE, colour_series=True):
    """openpyxl hides the axis numbers by default and puts the titles on top of
    them, so this switches the axes back on, keeps the titles outside the plot
    area and gives every chart the same colour"""
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.x_axis.majorTickMark = "out"
    chart.y_axis.majorTickMark = "out"
    chart.varyColors = False

    # overlay = False -> Excel reserves space for the title instead of
    # printing it over the axis numbers
    for title in (chart.title, chart.x_axis.title, chart.y_axis.title):
        if title is not None:
            title.overlay = False

    if colour_series:
        for s in chart.series:
            s.graphicalProperties = GraphicalProperties(solidFill=colour)
            s.graphicalProperties.line = LineProperties(solidFill=colour)


# ---------- CHART 1 : bar chart - planets by status ----------
bar1 = BarChart()
bar1.type = "bar"                      # horizontal bars
bar1.title = "Number of Planets by Planet Status"
# in a horizontal bar chart Excel draws the category axis (x_axis) on the
# left side and the value axis (y_axis) at the bottom
bar1.x_axis.title = "Planet Status"
bar1.y_axis.title = "Number of Planets"
data = Reference(ws2, min_col=2, min_row=2, max_row=2 + len(status_tbl))
cats = Reference(ws2, min_col=1, min_row=3, max_row=2 + len(status_tbl))
bar1.add_data(data, titles_from_data=True)
bar1.set_categories(cats)
bar1.legend = None
bar1.height = 7.5
bar1.width = 11.94
bar1.gapWidth = 200
value_labels(bar1)
fix_chart(bar1)
ws3.add_chart(bar1, "B14")

# ---------- CHART 2 : column chart - top 10 planets by mass ----------
bar2 = BarChart()
bar2.type = "col"
bar2.title = "Top 10 Planets by Mass"
bar2.y_axis.title = "Mass (MJup)"
bar2.x_axis.title = "Planet Name"
first = end1 + 4                       # header row of table 2
data = Reference(ws2, min_col=2, min_row=first, max_row=first + 10)
cats = Reference(ws2, min_col=1, min_row=first + 1, max_row=first + 10)
bar2.add_data(data, titles_from_data=True)
bar2.set_categories(cats)
bar2.legend = None
bar2.height = 7.5
bar2.width = 15.92
bar2.y_axis.scaling.min = 0            # start from 0 so the bars are not misleading
bar2.y_axis.numFmt = "0.0"
# the top 10 masses are all close to each other, so the exact value is
# printed on top of every bar
value_labels(bar2, fmt="0.00")
fix_chart(bar2)
ws3.add_chart(bar2, "H14")

# ---------- CHART 3 : scatter plot - mass vs radius ----------
sc = ScatterChart()
sc.title = "Relationship between Planet Mass and Radius"
sc.x_axis.title = "Mass (MJup)"
sc.y_axis.title = "Radius (RJup)"
last = 2 + len(scatter_df)
xval = Reference(ws2, min_col=6, min_row=3, max_row=last)
yval = Reference(ws2, min_col=7, min_row=2, max_row=last)
ser = Series(yval, xval, title_from_data=True)
ser.marker = Marker(symbol="circle", size=4)
ser.marker.graphicalProperties = GraphicalProperties(solidFill=BLUE)
ser.marker.graphicalProperties.line = LineProperties(solidFill=BLUE)
ser.graphicalProperties = GraphicalProperties()
ser.graphicalProperties.line.noFill = True        # points only, no line
sc.series.append(ser)
sc.legend = None
sc.height = 8.5
sc.width = 15.92
sc.x_axis.scaling.min = 0              # mass and radius can never be negative
sc.x_axis.scaling.max = 13
sc.y_axis.scaling.min = 0
sc.y_axis.scaling.max = 3
sc.x_axis.numFmt = "0.0"
sc.y_axis.numFmt = "0.0"
fix_chart(sc, colour_series=False)
ws3.add_chart(sc, "B30")

# ---------- CHART 4 : column chart - planets by mass category ----------
bar3 = BarChart()
bar3.type = "col"
bar3.title = "Number of Planets by Mass Category"
bar3.y_axis.title = "Number of Planets"
bar3.x_axis.title = "Mass Category"
first = end2 + 4
data = Reference(ws2, min_col=2, min_row=first, max_row=first + len(cat_tbl))
cats = Reference(ws2, min_col=1, min_row=first + 1, max_row=first + len(cat_tbl))
bar3.add_data(data, titles_from_data=True)
bar3.set_categories(cats)
bar3.legend = None
bar3.height = 8.5
bar3.width = 11.94
value_labels(bar3)
fix_chart(bar3)
ws3.add_chart(bar3, "J30")

# ---------- key findings box ----------
r = 48
ws3.merge_cells(start_row=r, start_column=2, end_row=r, end_column=15)
h = ws3.cell(row=r, column=2, value="KEY FINDINGS")
h.font = white_bold
h.fill = PatternFill("solid", fgColor=DARK)
h.alignment = Alignment(horizontal="left", vertical="center", indent=1)
for cc in range(2, 16):
    ws3.cell(row=r, column=cc).fill = PatternFill("solid", fgColor=DARK)

hi = df.loc[df["mass_mj"].idxmax()]
lo = df.loc[df["mass_mj"].idxmin()]
findings = [
    f"1. All {len(df):,} planets in the dataset have the status 'Confirmed', "
    f"so the status chart has only one bar.",
    f"2. Heaviest planet: {hi['name']} ({hi['mass_mj']:.3f} MJup). "
    f"Lightest planet: {lo['name']} ({lo['mass_mj']:.5f} MJup).",
    f"3. Average mass is {mass.mean():.3f} MJup and average radius is "
    f"{radius.mean():.3f} RJup, but the medians are much lower "
    f"({mass.median():.3f} and {radius.median():.3f}) - the data is right skewed.",
    "4. Mass and radius are positively related (Pearson r = 0.43, "
    "Spearman r = 0.82), but the relation is curved, not a straight line.",
    "5. Mass was available for only 51% of the planets, so the mass charts "
    "use 3,057 planets while the dataset has 5,986.",
]
for i, text in enumerate(findings):
    rr = r + 1 + i
    ws3.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=15)
    cell = ws3.cell(row=rr, column=2, value=text)
    cell.font = Font(size=10)
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for cc in range(2, 16):
        ws3.cell(row=rr, column=cc).fill = PatternFill("solid", fgColor=GREY)
    ws3.row_dimensions[rr].height = 16

# ---------- footer ----------
foot = r + len(findings) + 2
ws3.merge_cells(start_row=foot, start_column=2, end_row=foot, end_column=15)
f = ws3.cell(row=foot, column=2,
             value="Prepared by Dhruvan Karthik (2511021240019)")
f.font = Font(size=8, italic=True, color="808080")
f.alignment = Alignment(horizontal="right", vertical="center")

ws3.sheet_view.zoomScale = 90
wb.calculation.fullCalcOnLoad = True     # make Excel calculate the KPI formulas
wb.active = 2                            # open on the Dashboard sheet

# file properties (shown when you right click the file -> Details)
wb.properties.creator = "Dhruvan Karthik (2511021240019)"
wb.properties.lastModifiedBy = "Dhruvan Karthik"
wb.properties.title = "Exoplanet Data Analysis Dashboard"
wb.properties.description = ("Data cleaning and visualisation mini project - "
                             "exoplanet catalogue")

wb.save(OUT_FILE)
print("Dashboard saved as", OUT_FILE)
print("Sheets:", wb.sheetnames)
print("Scatter plot points:", len(scatter_df))
